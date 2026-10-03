"""Real-model, real-file, real-HTTP local acceptance. Never calls a live provider."""
import base64
import hashlib
import io
import json
import tempfile
import threading
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from openpyxl import Workbook
from docx import Document
from aidlp_judge.engine import Engine
from aidlp_judge.store import Store
from aidlp_judge.schemas import RunRequest
from aidlp_judge.jev_client import JevClient
from aidlp_judge.delivery import HttpDelivery
from aidlp_judge.scenarios import SCENARIOS

ROOT=Path(__file__).resolve().parents[2]
def main():
 received=[]
 class Sink(BaseHTTPRequestHandler):
  def log_message(self,*args):pass
  def do_POST(self):
   data=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
   received.append(data)
   result=json.dumps({'id':data['id'],'content_hash':hashlib.sha256(data['text'].encode()).hexdigest()}).encode()
   self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(result)
 server=ThreadingHTTPServer(('127.0.0.1',0),Sink);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
 worker=JevClient('http://127.0.0.1:8311',(ROOT/'.local/jev-state/api.token').read_text().strip());rows=[]
 fixtures=ROOT/'evidence/judge/fixtures';fixtures.mkdir(parents=True,exist_ok=True)
 try:
  with tempfile.TemporaryDirectory() as tmp:
   store=Store(Path(tmp)/'state.sqlite');engine=Engine(store,worker,HttpDelivery(f'http://127.0.0.1:{server.server_port}/receive','acceptance-only'))
   for scenario in SCENARIOS:
    for variant in ['protected','allowed']:
     text=scenario.protected_text if variant=='protected' else scenario.allowed_text
     payload={'scenario_id':scenario.id,'variant':variant}
     if scenario.id=='customer_upload':
      wb=Workbook();sheet=wb.active;sheet.title='Records'
      for line in text.splitlines():sheet.append([v.strip() for v in line.split('|')])
      content=io.BytesIO();wb.save(content);file=fixtures/f'{scenario.id}-{variant}.xlsx';file.write_bytes(content.getvalue())
      payload.update(filename=file.name,file_base64=base64.b64encode(content.getvalue()).decode())
     elif scenario.id in {'m365_grounding','drive_personal'}:
      doc=Document();doc.add_paragraph(text);content=io.BytesIO();doc.save(content);file=fixtures/f'{scenario.id}-{variant}.docx';file.write_bytes(content.getvalue())
      payload.update(filename=file.name,file_base64=base64.b64encode(content.getvalue()).decode())
     before=len(received);d=engine.run(RunRequest(**payload));count=len(received)-before
     row={'scenario':scenario.id,'variant':variant,'decision':d.model_dump(mode='json'),'http_receipts':count,'passed':False}
     if variant=='protected' and scenario.id!='cloud_summary':
      assert d.action in {'block','review'} and count==0
      if scenario.id in {'m365_grounding','rag_acl'}:assert not d.findings and d.extraction['skipped_by_authorization']
     elif scenario.id=='cloud_summary' and variant=='protected':
      assert d.action=='redact' and count==1
      raw=received[-1]['text']
      assert all(v not in raw for v in ['010-1234-5678','test.customer@example.test','900101-1234567'])
      assert '배송' in raw
      row['supported_identifiers_absent_at_receiver']=True
     elif d.action=='review':
      assert count==0 and d.approval_id
      store.approve(d.approval_id,True)
      approved=engine.run(RunRequest(**payload,approval_id=d.approval_id))
      assert approved.delivered and len(received)-before==1
      replay=engine.run(RunRequest(**payload,approval_id=d.approval_id))
      assert not replay.delivered and len(received)-before==1
      row['approved_decision']=approved.model_dump(mode='json');row['replay_blocked']=True
     else:assert d.delivered and count==1
     row['passed']=True;rows.append(row)
     print(scenario.id,variant,d.action,'PASS',flush=True)
   audit=json.dumps(store.events(),ensure_ascii=False)
   assert 'test.customer@example.test' not in audit and '김가람' not in audit
   (ROOT/'evidence/judge/acceptance.json').write_text(json.dumps({'scope':'controlled_local_http_receiver','inference':'separate_authenticated_http_api','live_provider_connected':False,'model':'decider-4b','tests':rows,'passed':len(rows),'raw_content_persisted':False},ensure_ascii=False,indent=2)+'\n')
 finally:worker.close();server.shutdown();server.server_close()
if __name__=='__main__':main()
