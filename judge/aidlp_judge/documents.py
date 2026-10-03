"""Bounded local extraction. Unsupported or incomplete inspection is explicit."""
import base64
import binascii
import csv
import hashlib
import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_BYTES=4*1024*1024
MAX_CHARS=60000
class ExtractionError(ValueError):pass

def _bounded(text):
    if len(text)>MAX_CHARS:raise ExtractionError('extracted_text_limit')
    if not text.strip():raise ExtractionError('no_inspectable_text')
    if '\x00' in text:raise ExtractionError('unsupported_binary_content')
    return text

def _zip_guard(data):
    try:
        z=zipfile.ZipFile(io.BytesIO(data));entries=z.infolist()
        if len(entries)>1500 or sum(e.file_size for e in entries)>20*1024*1024:raise ExtractionError('expanded_archive_limit')
        if any(e.flag_bits&1 for e in entries):raise ExtractionError('encrypted_document')
        if any('vbaproject' in e.filename.lower() or '/embeddings/' in e.filename.lower() for e in entries):raise ExtractionError('embedded_content_unsupported')
        if any(e.file_size>max(e.compress_size,1)*200 for e in entries):raise ExtractionError('archive_ratio_limit')
        return z
    except (zipfile.BadZipFile,RuntimeError):raise ExtractionError('invalid_document') from None

def extract(filename:str,encoded:str):
    try:data=base64.b64decode(encoded,validate=True)
    except (binascii.Error,ValueError):raise ExtractionError('invalid_base64') from None
    if len(data)>MAX_BYTES:raise ExtractionError('upload_size_limit')
    suffix=Path(filename).suffix.lower()
    if suffix in {'.txt','.md','.csv','.tsv','.json'}:
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError:raise ExtractionError('unsupported_text_encoding') from None
        if suffix in {'.csv','.tsv'}:
            rows=csv.reader(io.StringIO(text),delimiter='\t' if suffix=='.tsv' else ',')
            text='\n'.join(' | '.join(row) for row in rows)
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        with _zip_guard(data) as z:
            if any(e.startswith('xl/externalLinks/') for e in z.namelist()):raise ExtractionError('external_links_unsupported')
            if any(e.startswith('xl/media/') for e in z.namelist()):raise ExtractionError('document_images_require_ocr')
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        lines=[];chars=0;cells=0
        try:
            for sheet in wb:
                if sheet.max_row and sheet.max_row>10000:raise ExtractionError('spreadsheet_row_limit')
                lines.append('[Sheet: '+sheet.title+']')
                for row in sheet.iter_rows():
                    vals=[]
                    for cell in row:
                        cells+=1
                        if cells>50000:raise ExtractionError('spreadsheet_cell_limit')
                        if cell.data_type=='f':raise ExtractionError('spreadsheet_formula_unsupported')
                        if cell.value is not None:vals.append(str(cell.value))
                        else:vals.append('')
                    line=' | '.join(vals)
                    chars+=len(line)
                    if chars>MAX_CHARS:raise ExtractionError('extracted_text_limit')
                    if line.strip(' |'):lines.append(line)
        finally:wb.close()
        text='\n'.join(lines)
    elif suffix=='.docx':
        lines=[]
        with _zip_guard(data) as z:
            for name in z.namelist():
                if name.startswith('word/') and name.endswith('.xml') and any(t in name for t in ['document','header','footer','footnotes','endnotes','comments']):
                    xml=z.read(name)
                    if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:raise ExtractionError('xml_entities_unsupported')
                    root=ET.fromstring(xml)
                    for p in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                        line=''.join(t.text or '' for t in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
                        if line:lines.append(line)
            if any(n.startswith('word/media/') for n in z.namelist()):raise ExtractionError('document_images_require_ocr')
        text='\n'.join(lines)
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=True)
        if reader.is_encrypted:raise ExtractionError('encrypted_document')
        catalog=reader.trailer['/Root']
        if any(k in catalog for k in ['/AcroForm','/Names','/OpenAction','/AA']):raise ExtractionError('pdf_active_or_embedded_content_unsupported')
        if len(reader.pages)>40:raise ExtractionError('pdf_page_limit')
        texts=[]
        for page in reader.pages:
            if '/Annots' in page:raise ExtractionError('pdf_annotations_unsupported')
            if list(page.images):raise ExtractionError('pdf_images_require_ocr')
            value=page.extract_text() or ''
            if not value.strip():raise ExtractionError('pdf_page_without_text')
            texts.append(value)
        text='\n'.join(texts)
    else:raise ExtractionError('unsupported_document_format')
    text=_bounded(text)
    return text,dict(format=suffix.lstrip('.'),bytes=len(data),characters=len(text),sha256=hashlib.sha256(data).hexdigest(),complete=True)


def _extract_worker(pipe,filename,encoded):
    try:
        # Parser exceptions can contain document data; return only known reason codes.
        try:pipe.send(('ok',extract(filename,encoded)))
        except ExtractionError as exc:pipe.send(('error',str(exc)))
        except Exception:pipe.send(('error','document_extraction_failed'))
    finally:pipe.close()


def extract_bounded(filename,encoded,timeout=10):
    """Hard parser deadline; spawned worker is terminated on malformed input hangs."""
    import multiprocessing as mp
    context=mp.get_context('spawn');parent,child=context.Pipe(duplex=False)
    proc=context.Process(target=_extract_worker,args=(child,filename,encoded),daemon=True)
    proc.start();child.close()
    try:
        if not parent.poll(timeout):raise ExtractionError('document_extraction_timeout')
        kind,value=parent.recv()
        if kind!='ok':raise ExtractionError(value)
        return value
    except (EOFError,BrokenPipeError):raise ExtractionError('document_extraction_failed') from None
    finally:
        parent.close();proc.join(timeout=0.1)
        if proc.is_alive():proc.terminate();proc.join(timeout=1)
        if proc.is_alive():proc.kill();proc.join(timeout=1)
        proc.close()
